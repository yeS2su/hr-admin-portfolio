from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
import sys

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))

from workbook_import import ImportItem, build_preview, mark_duplicate_candidates
from generate_import_sql import write_import_sql


class WorkbookImportPreviewTest(unittest.TestCase):
    def make_workbook(self, path: Path) -> None:
        workbook = Workbook()
        for index in range(3):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = f'가상사용자{index + 1}'
            sheet.append(['', '연', '월', '주', '일', '분류', '업무 내용', '완료여부', '월별 업무'])
            sheet.append(['', 2026, 9, '9/21-25', '9/21', '가상분류', f'가상 일일업무 {index + 1}', True, '가상 월별업무'])
            sheet.append(['', 2026, 9, '9/21-25', None, '', f'가상 주간업무 {index + 1}', False, None])
        reference = workbook.create_sheet('가상 주간업무보고')
        reference['A1'] = '가상 참고자료'
        workbook.save(path)

    def test_builds_privacy_safe_counts_for_three_users(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / 'mock.xlsx'
            self.make_workbook(source)

            preview = build_preview(source)

            self.assertEqual(preview['sheet_count'], 4)
            self.assertEqual(len(preview['users']), 3)
            self.assertEqual(preview['users'][0]['daily'], 1)
            self.assertEqual(preview['users'][0]['weekly'], 1)
            self.assertEqual(preview['users'][0]['monthly'], 1)
            self.assertEqual(preview['users'][0]['completed'], 1)
            self.assertEqual(preview['errors']['total'], 0)
            self.assertEqual(preview['reference_sheets']['weekly_summary'], 1)
            serialized = str(preview)
            self.assertNotIn('가상 일일업무', serialized)
            self.assertNotIn('가상사용자', serialized)

    def test_reports_invalid_dates_without_copying_content(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / 'invalid.xlsx'
            self.make_workbook(source)
            from openpyxl import load_workbook
            workbook = load_workbook(source)
            workbook.worksheets[0]['E2'] = '2/31'
            workbook.worksheets[0]['G2'] = '출력 금지 문구'
            workbook.save(source)

            preview = build_preview(source)

            self.assertEqual(preview['errors']['invalid_date'], 1)
            self.assertNotIn('출력 금지 문구', str(preview))


    def test_understands_day_with_weekday_decimal_date_and_weekday_label(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / 'date-shapes.xlsx'
            self.make_workbook(source)
            from openpyxl import load_workbook
            workbook = load_workbook(source)
            workbook.worksheets[0]['E2'] = '21(월)'
            workbook.worksheets[1]['E2'] = 9.21
            workbook.worksheets[2]['E2'] = '월-금'
            workbook.save(source)

            preview = build_preview(source)

            self.assertEqual(preview['users'][0]['daily'], 1)
            self.assertEqual(preview['users'][1]['daily'], 1)
            self.assertEqual(preview['users'][2]['daily'], 0)
            self.assertEqual(preview['users'][2]['weekly'], 2)
            self.assertEqual(preview['errors']['total'], 0)
            self.assertEqual(preview['warnings']['unparsed_day_label'], 1)

    def test_preserves_cross_month_yearly_and_undated_items(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / 'periods.xlsx'
            self.make_workbook(source)
            from openpyxl import load_workbook
            workbook = load_workbook(source)
            first = workbook.worksheets[0]
            first.append(['', 2026, None, None, None, '', '가상 연간업무', False, None])
            second = workbook.worksheets[1]
            second.append(['', None, None, None, None, '', '가상 기간 미지정 업무', False, None])
            third = workbook.worksheets[2]
            third['C2'] = 4
            third['D2'] = '30~ 4/3'
            third['E2'] = 31.0
            workbook.save(source)

            preview = build_preview(source)

            self.assertEqual(preview['users'][0]['yearly'], 1)
            self.assertEqual(preview['users'][1]['undated'], 1)
            self.assertEqual(preview['users'][2]['daily'], 1)
            self.assertEqual(preview['errors']['invalid_date'], 0)
            self.assertEqual(preview['warnings']['missing_period'], 2)
            self.assertNotIn('가상 기간 미지정 업무', str(preview))

    def test_marks_repeated_items_for_review_without_deleting_them(self) -> None:
        original = ImportItem(
            owner_user_id='mock-user', cadence='daily', work_date='2026-09-21',
            year=2026, month=9, week_label='9/21-25', day_label='9/21',
            category='가상분류', content='가상 중복업무', is_completed=False,
            source_sheet_key='sheet_01', source_cell='G2',
        )
        repeated = ImportItem(**{**original.__dict__, 'source_cell': 'G3'})

        marked, count = mark_duplicate_candidates([original, repeated])

        self.assertEqual(count, 1)
        self.assertEqual(len(marked), 2)
        self.assertFalse(marked[0].needs_review)
        self.assertTrue(marked[1].needs_review)
    def test_generates_complete_temporary_sql_for_synthetic_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / 'mock.xlsx'
            output = Path(temp_dir) / 'import.sql'
            self.make_workbook(source)

            report = write_import_sql(source, output)
            sql = output.read_text(encoding='utf-8')

            self.assertEqual(report['work_entries'], 9)
            self.assertEqual(report['reference_records'], 1)
            self.assertEqual(report['hard_errors'], 0)
            self.assertIn('BEGIN TRANSACTION', sql)
            self.assertIn('reference_records', sql)
            self.assertIn('COMMIT', sql)
if __name__ == '__main__':
    unittest.main()


